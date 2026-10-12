"""Сборка отчёта МХТП (вариант 8): otchet_shablon.md -> report.md -> DOCX + PDF; ЗАЩИТА.md -> DOCX + PDF;
HYSYS_инструкция.md -> МХТП_HYSYS_инструкция.docx + PDF (рисунки — pr_flash.py).

    python build_report.py                  (всё; из папки src или из корня репозитория)
    python build_report.py defense hysys    (только ЗАЩИТА и инструкция HYSYS)

Шаги: 1) все расчёты mhtp_v8.run_all() и графики в fig/;
      2) подстановка чисел {{ключ}}, кода {{code:имена}}, вывода программы {{out:ключ}},
         нумерация таблиц {{tab:ключ}}, рисунков {{fig:ключ}}, источников {{ref:ключ}};
      3) _lib.md2docx.build_docx (титул, оглавление, OMML-формулы);
      4) локальная доводка DOCX: шрифт кода 9 пт, пустая строка после подписи рисунка;
      5) PDF-превью через LibreOffice с обновлённым оглавлением (lo_pdf.py);
      6) линтер: в обычном тексте нет «*», «^», «×», десятичный разделитель — запятая.
"""
from __future__ import annotations

import contextlib
import inspect
import io
import re
import sys
from pathlib import Path

import numpy as np

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
DISC = HERE.parent
OUT = DISC / "out"
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(ROOT))

import mhtp_v8 as m  # noqa: E402
from _lib.md2docx import TitleSpec, build_docx  # noqa: E402
from lo_pdf import docx_to_pdf_toc  # noqa: E402

REPORT_DOCX = OUT / "МХТП_Отчет_7сем_Суханов_вар8.docx"
DEFENSE_DOCX = OUT / "ЗАЩИТА.docx"
HYSYS_DOCX = OUT / "МХТП_HYSYS_инструкция.docx"
MINUS = "−"


# ---------------------------------------------------------------- форматирование чисел
def f(x, nd=2):
    """Число с десятичной запятой и знаком минус: f(-1.5, 1) -> '−1,5'."""
    s = f"{x:.{nd}f}"
    if s.startswith("-") and float(s) == 0:
        s = s[1:]
    return s.replace("-", MINUS).replace(".", ",")


def fe(x, nd=3):
    """Экспоненциальная запись для формулы: 1.375e4 -> '1{,}375 \\cdot 10^{4}' (без $)."""
    if x == 0:
        return "0"
    e = int(np.floor(np.log10(abs(x))))
    mant = x / 10 ** e
    if round(abs(mant), nd) >= 10:
        mant, e = mant / 10, e + 1
    ms = f"{mant:.{nd}f}".replace(".", "{,}").replace("-", "-")
    return ms if e == 0 else f"{ms} \\cdot 10^{{{e}}}"


def fm(x, nd=3):
    """То же в виде строчной формулы $...$."""
    return f"${fe(x, nd)}$"


def pct(x, nd=2):
    return f(100 * x, nd)


# ---------------------------------------------------------------- код и вывод
SRC_TEXT = (HERE / "mhtp_v8.py").read_text(encoding="utf-8")


def _const_source(name):
    """Исходный текст присваивания NAME = ... (до баланса скобок)."""
    lines = SRC_TEXT.split("\n")
    for i, line in enumerate(lines):
        if re.match(rf"^{re.escape(name)}\s*=", line):
            j = i                                     # предшествующие строки-комментарии
            while j > 0 and lines[j - 1].startswith("#") and not lines[j - 1].startswith("# %%"):
                j -= 1
            out, depth = lines[j:i], 0
            for line2 in lines[i:]:
                out.append(line2)
                depth += sum(line2.count(c) for c in "([{") - sum(line2.count(c) for c in ")]}")
                if depth <= 0:
                    break
            return "\n".join(out)
    raise KeyError(name)


def code(names):
    parts = []
    for n in [s.strip() for s in names.split(",")]:
        obj = getattr(m, n)
        parts.append(inspect.getsource(obj).rstrip() if callable(obj) else _const_source(n))
    text = "\n\n".join(parts)
    long = [ln for ln in text.split("\n") if len(ln) > 86]
    if long:
        print(f"ВНИМАНИЕ: строки кода длиннее 86 символов в {names}:")
        for ln in long:
            print("   ", len(ln), ln)
    return "```python\n" + text + "\n```"


def capture(fn, *args):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        fn(*args)
    return "```\n" + buf.getvalue().rstrip() + "\n```"


# ---------------------------------------------------------------- значения для шаблона
def values(res):
    V, OUTS = {}, {}
    # ---- Т1 З1
    r = res["t1z1"]
    for k, c in enumerate(r["p"][::-1]):
        V[f"t11_a{k}"] = fe(c, 4)
    V["t11_a3abs"] = fe(abs(r["p"][1]), 4)
    V.update(t11_R2=f(r["R2"], 5), t11_SSE=f(r["SSE"], 2), t11_Ip=f(r["I_poly"], 0),
             t11_It=f(r["I_trap"], 0), t11_Cp=f(r["C_poly"], 2), t11_Ct=f(r["C_trap"], 2),
             t11_d=f(r["delta_rel"], 3), t11_dabs=f(abs(r["C_trap"] - r["C_poly"]), 3),
             t11_etr=f(r["err_trap_rule"], 3), t11_emod=f(r["err_model"], 3),
             t11_Cpm=f(r["C_poly_mass"], 3), t11_Ctm=f(r["C_trap_mass"], 3),
             t11_rmax=f(np.max(np.abs(r["resid"])), 2),
             t11_Tmax=f(m.T1_T[np.argmax(np.abs(r["resid"]))], 0))
    cp_ifp = lambda T: (-0.966 + 7.279e-2 * T - 3.755e-5 * T ** 2 + 7.58e-9 * T ** 3) * m.CAL
    rows = []
    for T, Ct in ((300, 73.93), (600, 128.7), (1000, 174.6)):
        rows.append(f"| {T} | {f(Ct, 2)} | {f(cp_ifp(T), 2)} | {f(Ct / m.M_C3H8, 3)} |")
    V["t11_ifp_rows"] = "\n".join(rows)
    rows = []
    n = len(m.T1_T)
    for i in range(0, n, 4):
        cells = []
        for j in range(i, i + 4):
            cells += [f(m.T1_T[j], 2).rstrip("0").rstrip(","), f(m.T1_C[j], 2)]
        rows.append("| " + " | ".join(cells) + " |")
    V["t11_data_rows"] = "\n".join(rows)
    OUTS["t1z1"] = capture(m.print_t1z1, r)
    # ---- Т1 З2
    r = res["t1z2"]
    V.update(t12_a0=f(r["a0"], 4), t12_b0=f(r["b0"], 4), t12_alin=f(r["a_lin"], 4),
             t12_blin=f(r["b_lin"], 4))
    fits = r["fits"]
    for k in fits:
        V[f"t12_{k}_R2"] = f(fits[k]["R2"], 4)
        V[f"t12_{k}_R2adj"] = f(fits[k]["R2adj"], 4)
        V[f"t12_{k}_SSE"] = f(fits[k]["SSE"], 5)
        V[f"t12_{k}_AIC"] = f(fits[k]["AIC"], 2)
        for j, (pv, pe) in enumerate(zip(fits[k]["p"], fits[k]["perr"])):
            V[f"t12_{k}_p{j}"] = f(pv, 4)
            V[f"t12_{k}_e{j}"] = f(pe, 4)
    V["t12_data_rows"] = "\n".join(f"| {f(x, 3)} | {f(y, 2)} |" for x, y in zip(m.X2, m.Y2))
    V["t12_res_rows"] = "\n".join(
        f"| {f(x, 3)} | {f(y, 2)} | " + " | ".join(f(fits[k]['yfit'][i], 3) for k in ("given", "power", "hyper", "exp")) + " |"
        for i, (x, y) in enumerate(zip(m.X2, m.Y2)))
    V["t12_maxres_given"] = f(np.max(np.abs(fits["given"]["resid"])), 3)
    V["t12_maxres_hyper"] = f(np.max(np.abs(fits["hyper"]["resid"])), 3)
    OUTS["t1z2"] = capture(m.print_t1z2, r)
    # ---- Т1 З3
    r = res["t1z3"]
    V.update(t13_Ea=f(r["Ea"] / 1000, 1), t13_Eaci=f(r["Ea_ci"] / 1000, 1), t13_A=fm(r["A"], 3),
             t13_lnA=f(np.log(r["A"]), 3), t13_lnAci=f(r["lnA_ci"], 3),
             t13_R2ln=f(r["R2_ln"], 4), t13_R2k=f(r["R2_k"], 4),
             t13_Eain=f(r["Ea_in"] / 1000, 1), t13_Ain=fm(r["A_in"], 3), t13_R2in=f(r["R2_in"], 4),
             t13_Eanl=f(r["Ea_nl"] / 1000, 1), t13_Anl=fm(r["A_nl"], 3), t13_R2nlk=f(r["R2_nl_k"], 4),
             t13_R2nlln=f(r["R2_nl_ln"], 2), t13_slope=f(-r["Ea"] / m.R, 1),
             t13_slopeabs=f(r["Ea"] / m.R, 1), t13_Am=fe(r["A"], 3),
             t13_r100=f(r["ratio"][0], 2), t13_r500=f(r["ratio"][-1], 2),
             t13_res150=f(r["resid_ln"][1], 3),
             t13_res500=f(r["resid_ln"][-1], 3),
             t13_step=f(m.T3_k[-1] / m.T3_k[-2], 2),
             t13_step_arr=f(np.exp(r["Ea"] / m.R * (1 / r["TK"][-2] - 1 / r["TK"][-1])), 2),
             t13_r150=f(r["ratio"][1], 2), t13_r400=f(r["ratio"][6], 2))
    V["t13_rows"] = "\n".join(
        f"| {f(t, 0)} | {f(T, 2)} | {f(1000 / T, 4)} | {f(k, 2)} | {f(np.log(k), 3)} | {f(kf, 3)} | {f(e, 3)} | {f(q, 2)} |"
        for t, T, k, kf, e, q in zip(m.T3_t, r["TK"], m.T3_k, r["k_fit"], r["resid_ln"], r["ratio"]))
    OUTS["t1z3"] = capture(m.print_t1z3, r)
    # ---- Т1 З4
    r = res["t1z4"]
    rows = []
    for deg, d in r["models"].items():
        coefs = "; ".join(f"$b_{j} = {fe(c, 4)}$" for j, c in enumerate(d["p"][::-1]))
        rows.append(f"| {deg} | {coefs} | {f(d['R2'], 5)} | {f(d['R2adj'], 5)} | {fm(d['SSE'], 3)} | "
                    f"{f(d['S'], 4)} | {fm(d['PRESS'], 3)} |")
        V[f"t14_R2_{deg}"] = f(d["R2"], 4)
        V[f"t14_SSE_{deg}"] = fm(d["SSE"], 2)
        V[f"t14_S_{deg}"] = f(d["S"], 4)
        V[f"t14_P_{deg}"] = fm(d["PRESS"], 2)
    V["t14_rows"] = "\n".join(rows)
    # таблица исходных данных: 4 строки × 3 пары «T – μ» (в одну строку 12 значений не помещаются)
    rows = ["| T, К | μ, мПа∙с | T, К | μ, мПа∙с | T, К | μ, мПа∙с |", "|" + ":---:|" * 6]
    for i in range(4):
        cells = []
        for j in (i, i + 4, i + 8):
            cells += [f(m.T4_T[j], 0), f(m.T4_mu[j], 3)]
        rows.append("| " + " | ".join(cells) + " |")
    V["t14_data_rows"] = "\n".join(rows)
    V.update(t14_F=f(r["F"], 1), t14_Fcrit=f(r["Fcrit"], 2), t14_R2and=f(r["R2_and"], 4),
             t14_SSEand=fm(r["SSE_and"], 2), t14_Aand=f(r["andrade"][0], 3), t14_Band=f(r["andrade"][1], 1),
             t14_mu260a=f(r["mu260_and"], 3), t14_mu260p=f(r["mu260_p3"], 3),
             t14_res260=f(r["models"][3]["resid"][1], 4))
    OUTS["t1z4"] = capture(m.print_t1z4, r)
    # ---- Т2
    for key in ("t2z1", "t2z2"):
        r = res[key]
        rows = []
        vid = r["res"]["ИГ"]["v_L"]
        names = {"ИГ": "Идеальный газ", "RK": "Редлиха–Квонга (RK)", "SRK": "Соаве–Редлиха–Квонга (SRK)",
                 "LK": "Ли–Кеслера (контроль)"}
        for k, d in r["res"].items():
            rows.append(f"| {names[k]} | {f(d['v_L'], 4)} | {fm(d['v'], 4)} | {f(d['Z'], 4)} | "
                        f"{f((d['v_L'] / vid - 1) * 100, 1)} |")
            kk = "IG" if k == "ИГ" else k
            V[f"{key}_{kk}_v"] = f(d["v_L"], 4)
            V[f"{key}_{kk}_Z"] = f(d["Z"], 4)
            V[f"{key}_{kk}_dev"] = f((d["v_L"] / vid - 1) * 100, 1)
        V[f"{key}_rows"] = "\n".join(rows)
        for eos in ("RK", "SRK"):
            d = r["res"][eos]
            V[f"{key}_{eos}_a"] = f(d["a"], 5)
            V[f"{key}_{eos}_b"] = fm(d["b"], 4)
            cr = [x for x in d["roots"] if abs(x.imag) > 1e-12]
            V[f"{key}_{eos}_cplx"] = (f"${fe(cr[0].real * 1000, 3)} \\pm {fe(abs(cr[0].imag) * 1000, 3)}i$"
                                      if cr else "—")
        V[f"{key}_Tr"] = f(r.get("Tr", r.get("Tpr", 0)), 4)
        V[f"{key}_Pr"] = f(r.get("Pr", r.get("Ppr", 0)), 4)
        V[f"{key}_P"] = f(r["P"] / m.MPA, 4)
        V[f"{key}_LKz0"] = f(r["res"]["LK"]["z0"], 4)
        V[f"{key}_LKz1"] = f(r["res"]["LK"]["zcorr"], 4)
    r1 = res["t2z1"]["res"]
    V.update(t2z1_over_RK=f((r1["ИГ"]["v"] / r1["RK"]["v"] - 1) * 100, 1),
             t2z1_over_SRK=f((r1["ИГ"]["v"] / r1["SRK"]["v"] - 1) * 100, 1),
             t2z1_SRKvsRK=f((r1["SRK"]["v"] / r1["RK"]["v"] - 1) * 100, 1),
             t2z1_LKvsSRK=f(abs(r1["LK"]["v"] / r1["SRK"]["v"] - 1) * 100, 1),
             t2z1_dRK=f(abs(r1["RK"]["v"] / r1["ИГ"]["v"] - 1) * 100, 1),
             t2z1_dSRK=f(abs(r1["SRK"]["v"] / r1["ИГ"]["v"] - 1) * 100, 1))
    r2_ = res["t2z2"]["res"]
    V.update(t2z2_aratio=f((1 - r2_["SRK"]["a"] / r2_["RK"]["a"]) * 100, 0),
             t2z2_maxdev=f(max(abs(r2_[k]["v"] / r2_["ИГ"]["v"] - 1) for k in ("RK", "SRK", "LK")) * 100, 1))
    r = res["t2z2"]
    V.update(t2z2_Tpc=f(r["Tpc"], 2), t2z2_Ppc=f(r["Ppc"], 2), t2z2_w=f(r["w"], 4))
    rows = []
    for i, nme in enumerate(r["names"]):
        rows.append(f"| {nme.replace('2', '~2~')} | {f(r['y'][i], 4)} | {f(r['res']['RK']['ai'][i], 5)} | "
                    f"{f(r['res']['SRK']['ai'][i], 5)} | {fm(r['res']['RK']['bi'][i], 4)} | {f(r['T'] / m.CRIT[nme][0], 3)} |")
    rows.append(f"| Смесь | {f(1, 4)} | {f(r['res']['RK']['a'], 5)} | {f(r['res']['SRK']['a'], 5)} | "
                f"{fm(r['res']['RK']['b'], 4)} | — |")
    V["t2z2_ab_rows"] = "\n".join(rows)
    OUTS["t2z1"] = capture(m.print_t2, res["t2z1"], "Аммиак:")
    OUTS["t2z2"] = capture(m.print_t2, res["t2z2"], "Смесь CO2/O2/N2:")
    # ---- Т3 З1
    r = res["t3z1"]
    V.update(t31_tb=f(r["tb"], 1), t31_td=f(r["td"], 1), t31_Tb=f(r["Tb"], 2), t31_Td=f(r["Td"], 2),
             t31_Ppsia=f(r["P_psia"], 2), t31_P=f(r["P"] / m.MPA, 4),
             t31_TbR=f(r["Tb"] * 1.8, 1), t31_TdR=f(r["Td"] * 1.8, 1),
             t31_Kratio=f(r["Kb"][0] / r["Kb"][-1], 0), t31_range=f(r["td"] - r["tb"], 0),
             t31_yCH4=pct(r["y_bubble"][0], 0), t31_xC5=pct(r["x_dew"][-1], 0))
    nm = {"CH4": "Метан", "C2H6": "Этан", "C3H8": "Пропан", "n-C5H12": "н-Пентан",
          "i-C4H10": "Изобутан", "n-C4H10": "н-Бутан"}
    V["t31_rows"] = "\n".join(
        f"| {nm[n]} | {f(z, 2)} | {f(kb, 4)} | {f(yb, 4)} | {f(kd, 4)} | {f(xd, 4)} |"
        for n, z, kb, yb, kd, xd in zip(m.T3Z1_NAMES, m.T3Z1_Z, r["Kb"], r["y_bubble"], r["Kd"], r["x_dew"]))
    V["t31_rows"] += f"\n| Сумма | {f(1, 2)} | — | {f(np.sum(r['y_bubble']), 4)} | — | {f(np.sum(r['x_dew']), 4)} |"
    V["mcw_rows"] = "\n".join(
        f"| {nm[n]} | {f(m.MCW[n][0], 1)} | {f(m.MCW[n][1], 5)} | {f(m.MCW[n][2], 5)} | {f(m.MCW[n][3], 5)} | ({m.MCW[n][4]}) |"
        for n in m.SEP_NAMES)
    OUTS["t3z1"] = capture(m.print_t3z1, r)
    # ---- Т3 З2
    b = res["t3z2_base"]
    tr = b["train"]
    fl = {"CH4": "CH~4~", "C2H6": "C~2~H~6~", "C3H8": "C~3~H~8~", "i-C4H10": "i-C~4~H~10~",
          "n-C4H10": "н-C~4~H~10~", "n-C5H12": "н-C~5~H~12~"}
    rows = []
    for i, n in enumerate(m.SEP_NAMES):
        cells = [fl[n], f(m.SEP_Z[i], 5)]
        for s in tr:
            cells += [f(s["x"][i], 5), f(s["y"][i], 5)]
        rows.append("| " + " | ".join(cells) + " |")
    cells = ["Всего", f(m.SEP_Z.sum(), 5)]
    for s in tr:
        cells += [f(np.sum(s["x"]), 5), f(np.sum(s["y"]), 5)]
    rows.append("| " + " | ".join(cells) + " |")
    cells = ["Доля отгона e", "—"]
    for s in tr:
        cells += [f(s["e"], 5), ""]
    rows.append("| " + " | ".join(cells) + " |")
    V["t32_base_rows"] = "\n".join(rows)
    V["t32_K_rows"] = "\n".join(
        f"| {fl[n]} | " + " | ".join(f(s["K"][i], 4) for s in tr) + " |" for i, n in enumerate(m.SEP_NAMES))
    V["t32_flow_rows"] = "\n".join(
        f"| С{j+1} | {f(s['P']/m.MPA, 1)} | {f(s['T']-273.15, 1)} | {f(s['P']/m.PSI, 1)} | {f(s['T']*1.8, 1)} | "
        f"{f(s['F'], 4)} | {f(s['V'], 4)} | {f(s['L'], 4)} | {fm(s['rr'], 1)} |" for j, s in enumerate(tr))
    for j, s in enumerate(tr, 1):
        V[f"t32_e{j}"] = f(s["e"], 4)
        V[f"t32_V{j}"] = f(s["V"], 4)
        V[f"t32_L{j}"] = f(s["L"], 4)
        V[f"t32_KCH4_{j}"] = f(s["K"][0], 3)
        V[f"t32_KC3_{j}"] = f(s["K"][2], 4)
    sb = b["summary"]
    V.update(t32_Vout=f(sb["V_out"], 4), t32_recC3=pct(sb["rec_C3"], 1), t32_lossCH4=pct(sb["loss_CH4"], 1),
             t32_yC3p=pct(sb["yC3plus"], 2), t32_sharp=f(sb["sharp"], 3), t32_yCH4out=pct(sb["y_out"][0], 2),
             t32_recC4=pct(sb["rec_C4plus"], 2), t32_yC3out=pct(sb["y_out"][2], 3),
             t32_xCH4_3=pct(tr[2]["x"][0], 1), t32_L1pct=pct(tr[0]["L"], 0),
             t32_c2pct=pct(1 - tr[1]["e"], 0), t32_c3pct=pct(1 - tr[2]["e"], 0))
    st = res["t3z2_study"]
    V["t32_P2dew"] = f(st["P2_dew"], 2)
    rows = []
    for s in st["rows"]:
        if abs(s["P2"] * 2 - round(s["P2"] * 2)) < 1e-9 or abs(s["P2"] - 9.8) < 1e-9 or abs(s["P2"] - 8.7) < 1e-9:
            st2 = []
            for e, stt, ef in zip(s["e"], s["states"], s["e_formal"]):
                if stt == "двухфазная":
                    st2.append(f(e, 4))
                elif stt == "газ":
                    st2.append(f"газ ({f(ef, 3) if np.isfinite(ef) else '∞'})")
                else:
                    st2.append(f"жидк. ({f(ef, 3) if np.isfinite(ef) else '−∞'})")
            gas = "нет газа" if s["V_out"] == 0 else f(s["yC3plus"] * 100, 3)
            mark = " (база)" if abs(s["P2"] - 8.7) < 1e-9 else (" (выбор)" if abs(s["P2"] - m.P2_CHOICE) < 1e-9 else "")
            rows.append(f"| {f(s['P2'], 1)}{mark} | " + " | ".join(st2[1:]) + f" | {f(s['V_out'], 4)} | {gas} | "
                        f"{pct(s['rec_C3'], 2)} | {pct(s['loss_CH4'], 2)} | {f(s['sharp'], 3)} |")
    V["t32_study_rows"] = "\n".join(rows)
    rws = {round(s["P2"], 1): s for s in st["rows"]}
    ch, ba = rws[m.P2_CHOICE], rws[8.7]
    V["P2_choice"] = f(m.P2_CHOICE, 1)
    comp_rows = []
    for lab, key, scale, nd in (("Доля отгона в С2, моль/моль", "e2", 1, 4), ("Доля отгона в С3, моль/моль", "e3", 1, 4),
                                ("Выход газа С3, моль/моль сырья", "V_out", 1, 4),
                                ("Содержание C~3+~ в газе С3, % мол.", "yC3plus", 100, 3),
                                ("Извлечение C~3~H~8~ в жидкость, %", "rec_C3", 100, 2),
                                ("Доля CH~4~ сырья, ушедшая в жидкость, %", "loss_CH4", 100, 2),
                                ("Показатель чёткости разделения Ч", "sharp", 1, 3)):
        g = (lambda s: s["e"][1]) if key == "e2" else (lambda s: s["e"][2]) if key == "e3" else (lambda s: s[key])
        comp_rows.append(f"| {lab} | {f(g(ba) * scale, nd)} | {f(g(ch) * scale, nd)} | {f((g(ch) - g(ba)) * scale, nd)} |")
    V["t32_cmp_rows"] = "\n".join(comp_rows)
    V.update(t32c_e2=f(ch["e"][1], 4), t32c_V=f(ch["V_out"], 4), t32c_y=f(ch["yC3plus"] * 100, 2),
             t32c_rec=pct(ch["rec_C3"], 1), t32c_loss=pct(ch["loss_CH4"], 1), t32c_sharp=f(ch["sharp"], 3),
             t32b_sharp=f(ba["sharp"], 3), t32_gain=f((ch["V_out"] / ba["V_out"] - 1) * 100, 0),
             t32_r40_y=f(rws[4.0]["yC3plus"] * 100, 2), t32_r40_rec=pct(rws[4.0]["rec_C3"], 1),
             t32_r40_loss=pct(rws[4.0]["loss_CH4"], 1), t32_r95_e3=f(rws[9.5]["e"][2], 3),
             t32_r95_loss=pct(rws[9.5]["loss_CH4"], 1), t32_r75_y=f(rws[7.5]["yC3plus"] * 100, 2),
             t32_r75_sharp=f(rws[7.5]["sharp"], 3), t32_r55_e2=f(rws[5.5]["e"][1], 4),
             t32_r98_e2=f(rws[9.8]["e"][1], 4), t32_r98_ef=f(rws[9.8]["e_formal"][2], 3))
    alt = m.t3z2_study(P2_grid=[9.0, 9.5, 9.8], rule="P3<=8.7")
    V["t32_alt"] = "; ".join(f"P~2~ = {f(s['P2'], 1)} МПа: выход газа {f(s['V_out'], 4)}, "
                             f"Ч = {f(s['sharp'], 3)}" for s in alt["rows"])
    OUTS["t3z2"] = capture(m.print_t3z2, b, st)
    # ---- Т4 З1
    r = res["t4z1"]
    nmz = ["N~2~", "H~2~", "NH~3~"]
    V["t41_rows"] = "\n".join(f"| {nmz[i]} | {f(r['n0'][i], 1)} | {f(r['n'][i], 4)} | {f(r['y'][i], 4)} |" for i in range(3))
    V["t41_rows"] += f"\n| Всего | {f(r['n0'].sum(), 1)} | {f(r['n'].sum(), 4)} | {f(1, 4)} |"
    V.update(t41_xi=f(r["xi"], 4), t41_cN2=pct(r["conv_N2"], 1), t41_cH2=pct(r["conv_H2"], 1),
             t41_KP=f(r["Kp"] * r["P_atm"], 0), t41_ntot=f(r["n"].sum(), 4), t41_nN2=f(r["n"][0], 4),
             t41_nH2=f(r["n"][1], 4), t41_nNH3=f(r["n"][2], 4), t41_yNH3=f(r["y"][2], 4))
    OUTS["t4z1"] = capture(m.print_t4z1, r)
    # ---- Т4 З2а
    r = res["t4z2a"]
    ke = r["ke"]
    rows = []
    for i, T in enumerate(m.DGF["T"]):
        rows.append(f"| {f(T, 0)} | {f(m.DHF['C2H6'][i], 0)} | {f(m.DHF['C2H4'][i], 0)} | {f(m.DHF['C2H2'][i], 0)} | "
                    f"{f(m.DGF['C2H6'][i], 0)} | {f(m.DGF['C2H4'][i], 0)} | {f(m.DGF['C2H2'][i], 0)} |")
    V["t42_ref_rows"] = "\n".join(rows)
    V["t42_dG_rows"] = "\n".join(
        f"| {f(T, 0)} | {f(ke['dG'][i], 0)} | {f(ke['dG'][i] * m.CAL / 1000, 2)} | {fm(np.exp(-ke['dG'][i] * m.CAL / (m.R * T)), 3)} |"
        for i, T in enumerate(m.DGF["T"]))
    V.update(t42_dG800=f(ke["dG"][-2], 0), t42_dG1000=f(ke["dG"][-1], 0), t42_dGlin=f(ke["dG_lin"], 1),
             t42_dGhs=f(ke["dG_hs"], 1), t42_dH=f(ke["dH"][-1], 0), t42_dS=f(ke["dS"], 3),
             t42_Klin=f(ke["K_lin"], 4), t42_Khs=f(ke["K_hs"], 4), t42_K1000=f(r["K1000"], 4),
             t42_dGlinJ=f(ke["dG_lin"] * m.CAL, 0), t42_x=f(r["x"], 4), t42_dxdK=f(r["dxdK"], 4),
             t42_K1_1367=f(r["K1_1367"][1], 1), t42_K2_1367=f(r["K2_1367"][1], 3),
             t42_K1_1367l=f(r["K1_1367"][0], 1), t42_K2_1367l=f(r["K2_1367"][0], 3),
             t42_slope=f((ke["dG"][-1] - ke["dG"][-2]) / 200, 3), t42_yC2H6=pct(r["y"][0], 1),
             t42_yC2H4=pct(r["y"][1], 1), t42_Kgrow=f(ke["K_lin"] / r["K1000"], 1),
             t42_Kdiff=f((1.76 / ke["K_lin"] - 1) * 100, 0),
             t42_dx=f(r["sens"]["Finlayson"][1] - r["sens"]["lin"][1], 3),
             t42_KK=f(abs(ke["K_hs"] / ke["K_lin"] - 1) * 100, 1))
    nm2 = ["C~2~H~6~", "C~2~H~4~", "H~2~"]
    V["t42_rows"] = "\n".join(f"| {nm2[i]} | {f([1, 0, 0][i], 1)} | {f(r['n'][i], 4)} | {f(r['y'][i], 4)} |" for i in range(3))
    V["t42_rows"] += f"\n| Всего | {f(1, 1)} | {f(r['n'].sum(), 4)} | {f(1, 4)} |"
    lab = {"lin": "Линейная экстраполяция ΔG°r (основной)", "hs": "ΔH°r, ΔS°r = const (от 1000 К)",
           "Finlayson": "Значение из задачника [{{ref:finlayson}}, задача 4.12]"}
    V["t42_sens_rows"] = "\n".join(
        f"| {lab[k]} | {f(Kv, 4)} | {f(x, 4)} | {f((1 - x) / (1 + x), 4)} | {f(x / (1 + x), 4)} |"
        for k, (Kv, x) in r["sens"].items())
    OUTS["t4z2a"] = capture(m.print_t4z2a, r)
    # ---- Т4 З2б
    r = res["t4z2b"]
    nm3 = ["C~2~H~6~", "C~2~H~4~", "H~2~", "C~2~H~2~"]
    V["t43_rows"] = "\n".join(f"| {nm3[i]} | {f([1, 0, 0, 0][i], 1)} | {f(r['n'][i], 4)} | {f(r['y'][i], 4)} |"
                              for i in range(4))
    V["t43_rows"] += f"\n| Всего | {f(1, 1)} | {f(r['n'].sum(), 4)} | {f(1, 4)} |"
    V.update(t43_x1=f(r["x1"], 4), t43_x2=f(r["x2"], 4), t43_res=fm(np.max(np.abs(r["residual"])), 0),
             t43_raw1=f(r["fsolve_raw"][0], 1), t43_raw2=f(r["fsolve_raw"][1], 1),
             t43_sel=pct(r["n"][3] / (1 - r["n"][0]), 1), t43_yC2H6=pct(r["y"][0], 1),
             t43_yC2H4=pct(r["y"][1], 1), t43_yH2=pct(r["y"][2], 1), t43_yC2H2=pct(r["y"][3], 1))
    OUTS["t4z2b"] = capture(m.print_t4z2b, r)
    # ---- Т5 З1
    r = res["t5z1"]
    order = ["F", "D", "B", "D1", "B1", "D2", "B2"]
    rows = []
    for i, c in enumerate(m.T5_COMP):
        rows.append(f"| {c} | " + " | ".join(f(r["flows"][s] * r["comp"][s][i], 2) for s in order) + " |")
    rows.append("| Всего | " + " | ".join(f(r["flows"][s], 2) for s in order) + " |")
    rows.append("| **Доля** | | | | | | | |")
    for i, c in enumerate(m.T5_COMP):
        rows.append(f"| {c} | " + " | ".join(f(r["comp"][s][i], 4) for s in order) + " |")
    rows.append("| Всего | " + " | ".join(f(r["comp"][s].sum(), 4) for s in order) + " |")
    V["t51_rows"] = "\n".join(rows)
    for i, c in enumerate(("X", "S", "T", "B")):
        V[f"t51_xD{c}"] = pct(r["comp"]["D"][i], 1)
        V[f"t51_xB{c}"] = pct(r["comp"]["B"][i], 1)
    U_, S_, Vt_ = np.linalg.svd(r["A"])
    vv = Vt_[-1] / -Vt_[-1][1]
    V.update(t51_cD2=f(vv[2], 3), t51_cB2=f(vv[3], 3))
    V.update(t51_cond=f(r["cond"], 1), t51_det=fm(r["det"], 3), t51_svmin=f(r["sv"][-1], 5),
             t51_svmax=f(r["sv"][0], 4))
    for d, mc in r["mc"].items():
        key = "01" if d == 0.001 else "025"
        V[f"t51_mc{key}"] = "; ".join(f"{nme} {f(lo, 0)}–{f(hi, 0)}" for nme, lo, hi in zip(("D~1~", "B~1~", "D~2~", "B~2~"), mc["p5"], mc["p95"]))
        V[f"t51_neg{key}"] = f(mc["neg"] * 100, 1)
    OUTS["t5z1"] = capture(m.print_t5z1, r)
    # ---- Т5 З2
    r = res["t5z2"]
    order = ["1", "2", "R", "3", "4", "5", "6", "7"]
    rows = []
    for i, c in enumerate(("Толуол", "Бензол", "Ксилол")):
        rows.append(f"| {c} | " + " | ".join(f(r["streams"][s][i], 2) for s in order) + " |")
    rows.append("| Всего | " + " | ".join(f(r["streams"][s].sum(), 2) for s in order) + " |")
    V["t52_rows"] = "\n".join(rows)
    V.update(t52_iters=str(r["iters"]), t52_xi=f(r["xi"], 4), t52_T7=f(r["streams"]["7"][0], 4),
             t52_T2=f(r["streams"]["2"][0], 4), t52_cX=pct(r["conv_to_xylene"], 2), t52_cB=pct(r["conv_to_benzene"], 2),
             t52_cT=pct(r["conv_total"], 2), t52_loss=pct(1 - r["conv_total"], 2), t52_yX=pct(r["yield_xylene"], 2),
             t52_rr=f(r["recycle_ratio"], 4), t52_min=f(r["mass_in"], 1),
             t52_X6=f(r["streams"]["6"][2], 4), t52_B4=f(r["streams"]["4"][1], 4),
             t52_F4=f(r["streams"]["4"].sum(), 2), t52_F6=f(r["streams"]["6"].sum(), 2),
             t52_pB4=pct(r["streams"]["4"][1] / r["streams"]["4"].sum(), 2),
             t52_pX6=pct(r["streams"]["6"][2] / r["streams"]["6"].sum(), 2))
    hist = r["hist"]
    V["t52_iter_rows"] = "\n".join(f"| {int(h[0])} | {f(h[1], 4)} | {f(h[2], 4)} | {f(h[3], 4)} |"
                                   for h in hist if int(h[0]) in (1, 2, 3, 5, 10, 15, 20, 30) or int(h[0]) == r["iters"])
    OUTS["t5z2"] = capture(m.print_t5z2, r)
    return V, OUTS


# ---------------------------------------------------------------- шаблон
def render(template: str, V: dict, OUTS: dict) -> str:
    text = template
    # код и вывод программы
    text = re.sub(r"\{\{code:([^}]+)\}\}", lambda mo: code(mo.group(1)), text)
    text = re.sub(r"\{\{out:([^}]+)\}\}", lambda mo: OUTS[mo.group(1)], text)
    # значения (могут содержать ссылки {{ref:..}}, поэтому до нумерации)
    def val(mo):
        k = mo.group(1)
        if k not in V:
            raise KeyError(f"нет значения для {{{{{k}}}}}")
        return str(V[k])
    for _ in range(2):
        text = re.sub(r"\{\{([A-Za-z0-9_]+)\}\}", val, text)
    # нумерация таблиц и рисунков по порядку подписей
    for kind, pat in (("tab", r"^: Таблица \{\{tab:([A-Za-z0-9_]+)\}\}"), ("fig", r"!\[Рисунок \{\{fig:([A-Za-z0-9_]+)\}\}")):
        order = []
        for mo in re.finditer(pat, text, flags=re.M):
            if mo.group(1) in order:
                raise ValueError(f"повтор подписи {kind}:{mo.group(1)}")
            order.append(mo.group(1))
        num = {k: str(i + 1) for i, k in enumerate(order)}
        def rep(mo):
            if mo.group(1) not in num:
                raise KeyError(f"нет подписи для {kind}:{mo.group(1)}")
            return num[mo.group(1)]
        text = re.sub(r"\{\{" + kind + r"(?:ref)?:([A-Za-z0-9_]+)\}\}", rep, text)
    # источники — по порядку первого упоминания в тексте (до списка литературы)
    body, _, bib = text.partition("<!-- BIBLIOGRAPHY -->")
    order = []
    for mo in re.finditer(r"\{\{ref:([a-z0-9_]+)\}\}", body):
        if mo.group(1) not in order:
            order.append(mo.group(1))
    entries = dict(re.findall(r"^\[\[([a-z0-9_]+)\]\]\s*(.+)$", bib, flags=re.M))
    missing = [k for k in order if k not in entries]
    if missing:
        raise KeyError(f"нет описания источников: {missing}")
    unused = [k for k in entries if k not in order]
    if unused:
        print("ВНИМАНИЕ: источники без ссылок в тексте:", unused)
    num = {k: str(i + 1) for i, k in enumerate(order)}
    body = re.sub(r"\{\{ref:([a-z0-9_]+)\}\}", lambda mo: num[mo.group(1)], body)
    biblist = "\n".join(f"{i + 1}. {entries[k]}" for i, k in enumerate(order))
    # описания [[ключ]] — только источник данных; в документ идёт лишь нумерованный список
    text = body + "\n" + biblist + "\n" + bib.split("<!-- LIST -->", 1)[1]
    # контроль вложенных формул: «$$» допустим только как пара в выключной формуле
    for ln in text.split("\n"):
        if ln.count("$$") % 2 or (ln.count("$$") and not ln.lstrip().startswith("$$")):
            raise ValueError(f"вложенные $ в строке: {ln[:120]}")
    # десятичная запятая внутри формул: 13,80 -> 13{,}80 (иначе в OMML появляется пробел)
    def fix_math(mo):
        return re.sub(r"(\d),(\d)", r"\1{,}\2", mo.group(0))
    text = re.sub(r"\$\$.+?\$\$|\$[^$\n]+?\$", fix_math, text, flags=re.S)
    leftover = re.findall(r"\{\{[^}]*\}\}", text)
    if leftover:
        raise ValueError(f"неподставленные метки: {leftover[:5]}")
    return text


# ---------------------------------------------------------------- доводка DOCX и линтер
def tune_docx(path: Path, code_pt=9):
    import docx
    from docx.oxml import OxmlElement
    from docx.shared import Pt
    d = docx.Document(str(path))
    st = d.styles["Source Code"]
    st.font.size = Pt(code_pt)
    for name in ("Source Code",):
        for p in d.paragraphs:
            if p.style.name == name:
                for r in p.runs:
                    r.font.size = Pt(code_pt)
    # числа с десятичной запятой в формулах — «обычный текст» (m:nor): Word и LibreOffice
    # не добавляют пробел после запятой
    M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    for r in d.element.body.iter(M + "r"):
        t = r.find(M + "t")
        if t is None or not t.text:
            continue
        rpr0 = r.find(M + "rPr")
        sty = rpr0.find(M + "sty") if rpr0 is not None else None
        upright_word = (sty is not None and sty.get(M + "val") == "p"
                        and re.fullmatch(r"[A-Za-z]+", t.text))      # N, H, ln, exp, min...
        if re.fullmatch(r"[−-]?\d+,\d+", t.text) or upright_word:
            rpr = r.find(M + "rPr")
            if rpr is None:
                rpr = OxmlElement("m:rPr")
                r.insert(0, rpr)
            if rpr.find(M + "nor") is None:
                rpr.insert(0, OxmlElement("m:nor"))
    cap = d.styles["Image Caption"]
    cap.paragraph_format.space_after = Pt(0)
    for p in list(d.paragraphs):                      # пустая строка после подписи рисунка
        if p.style.name == "Image Caption":
            new = OxmlElement("w:p")
            p._p.addnext(new)
    d.save(str(path))


LINT_ALLOW = [r"\d{2}\.\d{2}\.\d{4}", r"https?://\S+", r"ISBN [\d-]+", r"Р 7\.0\.100",
              r"\bC4\.5\b", r"\bC5\.0\b",            # названия алгоритмов деревьев решений
              r"\b(mhtp_v8|pr_flash|check|build_report)\.py\b", r"\.hsc\b",
              # ссылки на задачи/пункты/разделы, в том числе перечнем: «задачи 1.1–1.4», «пунктах 3.1, 3.4 и 3.8»
              r"(задач\w*|пункт\w*|разд\w*|конспект\w*|сценари\w*)\s+\d+(\.\d+)+[аб]?"
              r"((\s*[–,]\s*|\s+и\s+)\d+(\.\d+)+[аб]?)*",
              r"^\s*\d+(\.\d+)+[аб]?(\s|$)",        # номер раздела/задания в начале строки
              r"(задач[аеиу]?|задани[еяю]|разд(ел|еле)|пункт[ае]?) \d+(\.\d+)+[аб]?"]


def lint_docx(path: Path) -> list[str]:
    import docx
    d = docx.Document(str(path))
    texts = []
    for p in d.paragraphs:
        if p.style.name in ("Source Code",) or p.style.name.lower().startswith(("heading", "toc")):
            continue
        texts.append(p.text)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                texts.extend(p.text for p in c.paragraphs)
    problems = []
    for t in texts:
        s = t
        for pat in LINT_ALLOW:
            s = re.sub(pat, "", s)
        for ch in ("*", "^", "×"):
            if ch in s:
                problems.append(f"символ «{ch}»: {t[:90]}")
        if "[[" in s or "]]" in s:
            problems.append(f"служебная метка [[...]]: {t[:90]}")
        if "¿" in s:
            problems.append(f"битый символ: {t[:90]}")
        if "{,}" in s or "\\" in s or "$" in s:
            problems.append(f"остаток разметки формул: {t[:90]}")
        if re.search(r"\d\.\d", s):
            problems.append(f"десятичная точка: {t[:90]}")
        if re.search(r"\d \- \d|[а-я] - [а-я]", s):
            problems.append(f"дефис вместо тире: {t[:90]}")
    return problems


def build_plain(md_path: Path, out: Path, label: str) -> list[str]:
    """Документ без титула: оглавление на первой странице, текст — со второй; DOCX + PDF + линтер."""
    zmd = "\\newpage\n\n" + md_path.read_text(encoding="utf-8")
    build_docx(zmd, out, title=None, toc=True, resource_dirs=[HERE, HERE / "fig"])
    tune_docx(out)
    probs = lint_docx(out)
    print(f"Линтер {label}:", "замечаний нет" if not probs else f"{len(probs)} замечаний")
    for p in probs[:40]:
        print("   ", p)
    print(docx_to_pdf_toc(out))
    return probs


def main(parts=("report", "defense", "hysys")):
    """parts — что собирать: report (отчёт), defense (ЗАЩИТА), hysys (инструкция HYSYS)."""
    probs = []
    if "report" in parts:
        res = m.run_all(plots=True, verbose=False)
        V, OUTS = values(res)
        tpl = (HERE / "otchet_shablon.md").read_text(encoding="utf-8")
        md = render(tpl, V, OUTS)
        (HERE / "report.md").write_text(md, encoding="utf-8")
        print("report.md записан:", len(md), "символов")
        title = TitleSpec(doc_type="ОТЧЁТ", subject="Моделирование химико-технологических процессов",
                          topic="Отчет по решению задач за 7 семестр по курсу Моделирование "
                                "химико-технологических процессов",
                          extra_lines=["Вариант 8"], checker="В.А. Горбунов")
        build_docx(HERE / "report.md", REPORT_DOCX, title=title, toc=True,
                   resource_dirs=[HERE, HERE / "fig"])
        tune_docx(REPORT_DOCX)
        p1 = lint_docx(REPORT_DOCX)
        print("Линтер отчёта:", "замечаний нет" if not p1 else f"{len(p1)} замечаний")
        for p in p1[:40]:
            print("   ", p)
        print(docx_to_pdf_toc(REPORT_DOCX))
        probs += p1
    if "defense" in parts:
        probs += build_plain(DISC / "ЗАЩИТА.md", DEFENSE_DOCX, "ЗАЩИТА")
    if "hysys" in parts:
        import pr_flash
        pr_flash.run_all(verbose=False, plots=True)            # рисунки fig/hysys_s1_*.png
        probs += build_plain(HERE / "HYSYS_инструкция.md", HYSYS_DOCX, "инструкции HYSYS")
    return 1 if probs else 0


if __name__ == "__main__":
    sys.exit(main(tuple(sys.argv[1:]) or ("report", "defense", "hysys")))
