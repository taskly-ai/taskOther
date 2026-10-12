"""Автоматизация Aspen HYSYS V15 для перекрёстной проверки МХТП (вариант 8).

Скрипт через COM-интерфейс Windows строит в HYSYS три сценария из
07_МХТП/out/МХТП_HYSYS_инструкция.docx, решает их, читает результаты, сравнивает
с отчётом и ориентиром Пенга–Робинсона и пишет всё в Excel.

    S1 – сепарационный блок (задача 3.2): базовый режим, режим 6,0 МПа и
         исследование P2 = P3 = 4,0…9,5 МПа;
    S2 – синтез аммиака в равновесном реакторе (задача 4.1);
    S3 – дегидрирование этана в реакторе Гиббса (задача 4.2а).

Запуск (Windows, HYSYS установлен, Python 3.9+, pip install pywin32 openpyxl):
    python hysys_auto.py diag              – проверить связь с HYSYS, записать hysys_diag.txt
    python hysys_auto.py all               – построить и посчитать S1, S2, S3
    python hysys_auto.py s1                – только сепараторы (s2, s3 – аналогично)
    python hysys_auto.py read s1 файл.hsc  – НЕ строить, а прочитать схему, собранную вручную
                                             по инструкции (имена потоков как в инструкции)
    python hysys_auto.py selftest          – прогон логики на эмуляторе (без HYSYS)

Результаты: hysys_results.xlsx и файлы S1_separators.hsc, S2_ammonia.hsc,
S3_gibbs.hsc рядом со скриптом; подробный журнал – hysys_log.txt.

Автопостроение через COM – экспериментальное: API HYSYS не проверялось на V15
у автора скрипта. Каждый шаг обёрнут в попытки нескольких вариантов вызова; если
какой-то шаг не удался, скрипт пишет, что сделать вручную, и предлагает режим
`read`. Журнал hysys_log.txt и hysys_diag.txt нужно прислать для доработки.
"""
from __future__ import annotations

import datetime as _dt
import os
import sys
import time
import traceback
from pathlib import Path

HERE = Path(__file__).resolve().parent
LOG = HERE / "hysys_log.txt"
XLSX = HERE / "hysys_results.xlsx"

# ----------------------------------------------------------------- данные варианта 8
S1_COMPONENTS = ["Methane", "Ethane", "Propane", "i-Butane", "n-Butane", "n-Pentane"]
S1_Z = [0.72211, 0.03608, 0.22363, 0.00646, 0.00466, 0.00706]
S1_FEED = dict(T_C=40.8, P_kPa=9800.0, F_kmolh=100.0)
S1_MODES = {  # режим: (T2 °C, P2 кПа, T3 °C, P3 кПа)
    "base": (12.8, 8700.0, -40.8, 8700.0),
    "6MPa": (12.8, 6000.0, -40.8, 6000.0),
}
S1_STUDY_P = [4000.0 + 500.0 * i for i in range(12)]  # 4000…9500 кПа

# Ожидаемые значения: отчёт (McWilliams) и ориентир PR (src/pr_flash.py, kij = 0)
S1_EXPECT = {
    "base": [
        ("С1: доля отгона e1", "e1", 0.7361, 1.0),
        ("С1: жидкость Liq1, kgmole/h", "Liq1", 26.39, 0.0),
        ("С2: доля отгона e2", "e2", 0.8761, 0.8021),
        ("С2: жидкость Liq2, kgmole/h", "Liq2", 9.12, 19.79),
        ("С3: доля отгона e3", "e3", 0.4775, None),
        ("Газ С3 (Gas3), kgmole/h", "Gas3", 30.80, None),
        ("CH4 в газе С3, мол. доля", "yCH4_Gas3", 0.99249, None),
    ],
    "6MPa": [
        ("С1: доля отгона e1", "e1", 0.7361, 1.0),
        ("С2: доля отгона e2", "e2", 0.9805, 0.8620),
        ("С3: доля отгона e3", "e3", 0.6932, 0.5412),
        ("Газ С3 (Gas3), kgmole/h", "Gas3", 50.03, 46.66),
        ("CH4 в газе С3, мол. доля", "yCH4_Gas3", 0.97979, 0.93664),
        ("C3H8 в газе С3, мол. доля", "yC3H8_Gas3", 0.00911, 0.04304),
    ],
}

S2_COMPONENTS = ["Nitrogen", "Hydrogen", "Ammonia"]
S2_FEED_FLOWS = [1.1, 3.0, 0.2]  # kgmole/h
S2_FEED = dict(T_C=315.85, P_kPa=220 * 101.325)
S2_KEQ = 0.05  # атм^-1, Basis = Partial Pressure
S2_EXPECT = [
    ("N2 на выходе, kgmole/h", "nN2", 0.3654),
    ("H2 на выходе, kgmole/h", "nH2", 0.7963),
    ("NH3 на выходе, kgmole/h", "nNH3", 1.6691),
    ("Всего Vap1, kgmole/h", "Ftot", 2.8309),
    ("y NH3", "yNH3", 0.5896),
    ("Liq1, kgmole/h", "Liq1", 0.0),
    ("Степень превращения N2, %", "XN2", 66.78),
    ("Глубина реакции, kgmole/h", "extent", 1.4691),
]

S3_COMPONENTS = ["Ethane", "Ethylene", "Hydrogen"]
S3_FEED = dict(T_C=826.85, P_kPa=101.325, F_kmolh=1.0)
S3_EXPECT = [
    ("Степень превращения этана x", "x", 0.7899),
    ("C2H6 в Vap1, мол. доля", "yC2H6", 0.1174),
    ("C2H4 в Vap1, мол. доля", "yC2H4", 0.4413),
    ("H2 в Vap1, мол. доля", "yH2", 0.4413),
    ("Vap1, kgmole/h", "Ftot", 1.7899),
    ("K = yC2H4∙yH2/yC2H6", "K", 1.659),
]

# Варианты имён типов операций HYSYS (TypeName); перебираются по очереди
OP_TYPES = {
    "separator": ["flashtank", "sep3op", "separatorop"],
    "cooler": ["coolerop", "cooler"],
    "equilibrium": ["eqmreactorop", "equilibriumreactorop", "eqreactor", "equilibriumreactor"],
    "gibbs": ["gibbsreactorop", "gibbsreactor"],
    "set": ["setop", "set"],
}


# ----------------------------------------------------------------- журнал
def log(msg: str = "", also_print: bool = True):
    line = msg
    if also_print:
        print(line)
    with open(LOG, "a", encoding="utf-8") as f:
        f.write(line + "\n")


class StepFailed(RuntimeError):
    """Шаг автопостроения не удался – нужен ручной шаг или режим read."""


def attempt(desc: str, *funcs):
    """Выполнить первый сработавший вариант вызова; иначе StepFailed."""
    errors = []
    for fn in funcs:
        try:
            res = fn()
            log(f"  ✓ {desc}")
            return res
        except Exception as e:  # noqa: BLE001 – COM бросает разные исключения
            errors.append(f"{getattr(fn, '__name__', 'lambda')}: {e!r}")
    log(f"  ✗ {desc}")
    for e in errors:
        log(f"      {e}", also_print=False)
    raise StepFailed(desc)


# ----------------------------------------------------------------- связь с HYSYS
def connect(dry: bool = False):
    if dry:
        import fake_hysys  # эмулятор для selftest
        return fake_hysys.Application(), "FAKE"
    try:
        import win32com.client as win32
    except ImportError:
        sys.exit("Нужен пакет pywin32: pip install pywin32 (только Windows).")
    last = None
    candidates = ["HYSYS.Application.V15.0", "HYSYS.Application.V15", "HYSYS.Application"]
    found = registry_progids()
    candidates += [p for p in found if p not in candidates]
    for progid in candidates:
        try:
            app = win32.Dispatch(progid)
            try:
                app.SimulationCases.Count  # это действительно HYSYS, а не посторонний COM-класс
            except Exception as e:  # noqa: BLE001
                last = RuntimeError(f"{progid}: не похоже на HYSYS ({e!r})")
                continue
            try:
                app.Visible = True
            except Exception:  # noqa: BLE001
                pass
            return app, progid
        except Exception as e:  # noqa: BLE001
            last = e
    msg = [f"Не удалось подключиться к HYSYS через COM: {last!r}"]
    if found:
        msg.append("COM-имена HYSYS в реестре: " + ", ".join(found) + " (все попробованы).")
    else:
        msg.append("В реестре Windows нет ни одного COM-имени HYSYS.*: COM-сервер HYSYS не зарегистрирован.")
    msg.append("Сначала добейтесь, чтобы HYSYS открывался ярлыком (см. README, раздел 4). Затем один раз в "
               "командной строке ОТ ИМЕНИ АДМИНИСТРАТОРА: \"<папка HYSYS>\\AspenHysys.exe\" /regserver "
               "и снова: python hysys_auto.py diag")
    (HERE / "hysys_diag.txt").write_text("\n".join(msg), encoding="utf-8")
    sys.exit("\n".join(msg))


def registry_progids():
    """COM-имена (ProgID) вида HYSYS.* / Aspen*HYSYS*, зарегистрированные в Windows."""
    try:
        import winreg
    except ImportError:
        return []
    out = []
    try:
        root = winreg.HKEY_CLASSES_ROOT
        i = 0
        while True:
            try:
                name = winreg.EnumKey(root, i)
            except OSError:
                break
            i += 1
            low = name.lower()
            # только сервер приложения HYSYS (HYSYS.Application[.Vxx]); прочие классы с «HYSYS»
            # в имени (например, AspenTech.Hydraulics…HYSYSEmulsionViscosityCorrelation) не годятся
            if low.startswith("hysys.application"):
                try:
                    winreg.CloseKey(winreg.OpenKey(root, name + "\\CLSID"))
                    out.append(name)
                except OSError:
                    pass
    except Exception:  # noqa: BLE001
        pass
    return sorted(out)


def new_case(app):
    case = attempt("новый файл моделирования",
                   lambda: app.SimulationCases.Add(),
                   lambda: app.SimulationCases.Add(""))
    try:
        case.Visible = True
    except Exception:  # noqa: BLE001
        pass
    return case


def open_case(app, path: Path):
    case = attempt(f"открыть {path.name}", lambda: app.SimulationCases.Open(str(path)))
    try:
        case.Visible = True
    except Exception:  # noqa: BLE001
        pass
    return case


# ----------------------------------------------------------------- базис
def setup_basis(case, components, fp_name):
    """Список компонентов + пакет Peng-Robinson; возвращает FluidPackage."""
    bm = case.BasisManager
    try:
        bm.StartBasisChange()
    except Exception:  # noqa: BLE001
        pass
    fp = attempt("пакет свойств",
                 lambda: bm.FluidPackages.Add(fp_name),
                 lambda: bm.FluidPackages.Add(fp_name, "Peng-Robinson"))
    attempt("Peng-Robinson",
            lambda: setattr(fp, "PropertyPackageName", "Peng-Robinson"),
            lambda: setattr(fp, "PropertyPackageName", "PengRob"),
            lambda: setattr(fp, "PropertyPackageName", "Peng Robinson"))

    def via_fp():
        for c in components:
            fp.Components.Add(c)

    def via_list():
        cl = bm.ComponentLists.Add(fp_name + "-CL")
        for c in components:
            cl.Components.Add(c)
        fp.ComponentList = cl

    def via_list_noname():
        cl = bm.ComponentLists.Add()
        for c in components:
            cl.Components.Add(c)
        fp.ComponentList = cl

    attempt("компоненты: " + ", ".join(components), via_fp, via_list, via_list_noname)
    try:
        bm.EndBasisChange()
    except Exception:  # noqa: BLE001
        pass
    attempt("пакет → схема",
            lambda: setattr(case.Flowsheet, "FluidPackage", fp),
            lambda: case.Flowsheet.SetFluidPackage(fp))
    return fp


def fp_component_names(fp, fallback):
    for get in (lambda: list(fp.Components.Names),
                lambda: [fp.Components.Item(i).Name for i in range(fp.Components.Count)]):
        try:
            names = [str(n) for n in get()]
            if names:
                return names
        except Exception:  # noqa: BLE001
            pass
    return list(fallback)


# ----------------------------------------------------------------- потоки и аппараты
def add_stream(fs, name):
    return attempt(f"поток {name}", lambda: fs.MaterialStreams.Add(name))


def add_energy(fs, name):
    return attempt(f"энергопоток {name}", lambda: fs.EnergyStreams.Add(name))


def set_var(stream, var, value, unit):
    """Задать T/P/расход с единицами; запасной путь – внутренние единицы HYSYS."""
    internal = {"C": value, "kPa": value, "kgmole/h": value / 3600.0}[unit]  # T °C, P кПа, F kgmole/s
    attempt(f"{stream_name(stream)}.{var} = {value} {unit}",
            lambda: getattr(stream, var).SetValue(value, unit),
            lambda: setattr(stream, var + "Value", internal))


def set_composition(stream, fractions):
    vals = tuple(float(x) for x in fractions)
    attempt(f"{stream_name(stream)}: состав",
            lambda: setattr(stream.ComponentMolarFraction, "Values", vals),
            lambda: setattr(stream, "ComponentMolarFractionValue", vals),
            lambda: stream.ComponentMolarFraction.SetValues(vals, ""))


def stream_name(s):
    try:
        return str(s.Name)
    except Exception:  # noqa: BLE001
        return "?"


def add_op(fs, name, kind):
    funcs = [(lambda t=t: fs.Operations.Add(name, t)) for t in OP_TYPES[kind]]
    return attempt(f"аппарат {name} ({kind})", *funcs)


def connect_separator(op, feed, vap, liq):
    attempt(f"{op.Name}: вход", lambda: op.FeedStreams.Add(feed), lambda: op.AttachFeed(feed))
    attempt(f"{op.Name}: газ", lambda: setattr(op, "VapourProduct", vap),
            lambda: setattr(op, "VapourProductStream", vap))
    attempt(f"{op.Name}: жидкость", lambda: setattr(op, "LiquidProduct", liq),
            lambda: setattr(op, "LiquidProductStream", liq))


def connect_cooler(op, feed, prod, q):
    attempt(f"{op.Name}: вход", lambda: setattr(op, "FeedStream", feed), lambda: op.FeedStreams.Add(feed))
    attempt(f"{op.Name}: выход", lambda: setattr(op, "ProductStream", prod),
            lambda: op.ProductStreams.Add(prod))
    attempt(f"{op.Name}: энергия", lambda: setattr(op, "EnergyStream", q))


def solve(case, timeout=120):
    try:
        case.Solver.CanSolve = True
    except Exception:  # noqa: BLE001
        pass
    t0 = time.time()
    while time.time() - t0 < timeout:
        try:
            if not case.Solver.IsSolving:
                return
        except Exception:  # noqa: BLE001
            return
        time.sleep(0.2)
    log("  ! расчёт не завершился за отведённое время")


def hold(case):
    try:
        case.Solver.CanSolve = False
    except Exception:  # noqa: BLE001
        pass


def save(case, name):
    path = HERE / name
    try:
        case.SaveAs(str(path))
        log(f"  ✓ сохранено: {path.name}")
    except Exception as e:  # noqa: BLE001
        log(f"  ! не удалось сохранить {path.name}: {e!r}")


# ----------------------------------------------------------------- чтение результатов
def get_stream(fs, name):
    for get in (lambda: fs.MaterialStreams.Item(name), lambda: fs.MaterialStreams(name)):
        try:
            return get()
        except Exception:  # noqa: BLE001
            pass
    raise StepFailed(f"нет потока «{name}» (проверьте имена как в инструкции)")


def read_var(stream, var, unit):
    """Прочитать переменную; None, если не рассчитана."""
    for get, k in ((lambda: getattr(stream, var).GetValue(unit), 1.0),
                   (lambda: getattr(stream, var + "Value"), 3600.0 if unit == "kgmole/h" else 1.0)):
        try:
            raw = float(get())
            if raw <= -32000 or raw != raw:  # «пусто» в HYSYS = −32767 (<empty>), NaN
                return None
            return raw * k
        except Exception:  # noqa: BLE001
            pass
    return None


def read_fractions(stream, names):
    for get in (lambda: stream.ComponentMolarFraction.GetValues(""),
                lambda: stream.ComponentMolarFractionValue,
                lambda: stream.ComponentMolarFraction.Values):
        try:
            vals = [float(x) for x in get()]
            if len(vals) == len(names) and all(v > -32000 for v in vals):
                return dict(zip(names, vals))
        except Exception:  # noqa: BLE001
            pass
    return {}


def vapour_fraction(stream):
    return read_var(stream, "VapourFraction", "")


def flow(stream):
    return read_var(stream, "MolarFlow", "kgmole/h")


# ----------------------------------------------------------------- сценарий 1
def s1_build(app):
    log("\n=== S1: сепарационный блок (задача 3.2) ===")
    case = new_case(app)
    hold(case)
    fp = setup_basis(case, S1_COMPONENTS, "PR-S1")
    names = fp_component_names(fp, S1_COMPONENTS)
    z = [S1_Z[S1_COMPONENTS.index(n)] if n in S1_COMPONENTS else 0.0 for n in names]
    fs = case.Flowsheet
    st = {n: add_stream(fs, n) for n in
          ["Feed", "Gas1", "Liq1", "Feed2", "Gas2", "Liq2", "Feed3", "Gas3", "Liq3"]}
    q1, q2 = add_energy(fs, "Q-100"), add_energy(fs, "Q-101")
    set_var(st["Feed"], "Temperature", S1_FEED["T_C"], "C")
    set_var(st["Feed"], "Pressure", S1_FEED["P_kPa"], "kPa")
    set_var(st["Feed"], "MolarFlow", S1_FEED["F_kmolh"], "kgmole/h")
    set_composition(st["Feed"], z)
    v100, v101, v102 = (add_op(fs, n, "separator") for n in ("V-100", "V-101", "V-102"))
    e100, e101 = add_op(fs, "E-100", "cooler"), add_op(fs, "E-101", "cooler")
    connect_separator(v100, st["Feed"], st["Gas1"], st["Liq1"])
    connect_cooler(e100, st["Gas1"], st["Feed2"], q1)
    connect_separator(v101, st["Feed2"], st["Gas2"], st["Liq2"])
    connect_cooler(e101, st["Gas2"], st["Feed3"], q2)
    connect_separator(v102, st["Feed3"], st["Gas3"], st["Liq3"])
    return case


def s1_set_mode(case, T2, P2, T3, P3):
    fs = case.Flowsheet
    hold(case)
    f2, f3 = get_stream(fs, "Feed2"), get_stream(fs, "Feed3")
    set_var(f2, "Temperature", T2, "C")
    set_var(f2, "Pressure", P2, "kPa")
    set_var(f3, "Temperature", T3, "C")
    set_var(f3, "Pressure", P3, "kPa")
    solve(case)


def s1_read(case):
    fs = case.Flowsheet
    s = {n: get_stream(fs, n) for n in ["Feed", "Liq1", "Feed2", "Liq2", "Feed3", "Gas3"]}
    names = list(S1_COMPONENTS)
    try:
        names = fp_component_names(case.Flowsheet.FluidPackage, S1_COMPONENTS)
    except Exception:  # noqa: BLE001
        pass
    y = read_fractions(s["Gas3"], names)
    return {
        "e1": vapour_fraction(s["Feed"]), "Liq1": flow(s["Liq1"]),
        "e2": vapour_fraction(s["Feed2"]), "Liq2": flow(s["Liq2"]),
        "e3": vapour_fraction(s["Feed3"]), "Gas3": flow(s["Gas3"]),
        "yCH4_Gas3": y.get("Methane"), "yC3H8_Gas3": y.get("Propane"),
    }


def s1_run(app, case=None, built=True):
    if case is None:
        case = s1_build(app)
    res = {}
    for mode, (T2, P2, T3, P3) in S1_MODES.items():
        log(f"-- режим {mode}: P2 = {P2:.0f} кПа, P3 = {P3:.0f} кПа")
        if built:
            s1_set_mode(case, T2, P2, T3, P3)
        else:
            solve(case)
        res[mode] = s1_read(case)
        if not built:
            log("   (режим read: читаю как есть; второй режим и исследование пропущены)")
            break
    study = []
    if built:
        log("-- исследование P2 = P3 = 4,0…9,5 МПа")
        for P in S1_STUDY_P:
            s1_set_mode(case, 12.8, P, -40.8, P)
            r = s1_read(case)
            study.append((P, r))
            log(f"   P = {P/1000:4.1f} МПа: e2 = {fmt(r['e2'])}, e3 = {fmt(r['e3'])}, "
                f"Gas3 = {fmt(r['Gas3'])} kgmole/h")
        s1_set_mode(case, *S1_MODES["6MPa"])
        save(case, "S1_separators.hsc")
    return res, study


# ----------------------------------------------------------------- сценарий 2
def s2_reaction(case, fp):
    """Набор реакций с одной равновесной реакцией (Fixed Keq, Partial Pressure, atm)."""
    bm = case.BasisManager

    def via_rpm():
        rpm = bm.ReactionPackageManager
        rset = rpm.ReactionSets.Add("Set-1")
        rxn = rpm.Reactions.Add("Rxn-1", "Equilibrium")
        return rset, rxn

    def via_rpm2():
        rpm = bm.ReactionPackageManager
        rxn = rpm.Reactions.Add("Rxn-1", "EquilibriumReaction")
        rset = rpm.ReactionSets.Add("Set-1")
        return rset, rxn

    rset, rxn = attempt("набор реакций Set-1 и реакция Rxn-1", via_rpm, via_rpm2)
    coeffs = {"Nitrogen": -0.5, "Hydrogen": -1.5, "Ammonia": 1.0}

    def stoich_a():
        for c, nu in coeffs.items():
            rxn.AddComponent(c, nu)

    def stoich_b():
        names = [str(n) for n in rxn.Components.Names]
        rxn.StoichCoeffs = tuple(coeffs.get(n, 0.0) for n in names)

    attempt("стехиометрия −0,5 N2 −1,5 H2 +1 NH3", stoich_a, stoich_b)
    for desc, funcs in (
        ("Basis = Partial Pressure", [lambda: setattr(rxn, "Basis", "Partial Pressure"),
                                      lambda: setattr(rxn, "BasisName", "Partial Pressure")]),
        ("Basis Units = atm", [lambda: setattr(rxn, "BasisUnits", "atm")]),
        ("Keq Source = Fixed Keq", [lambda: setattr(rxn, "KeqSource", "Fixed Keq"),
                                    lambda: setattr(rxn, "EquilibriumConstantSource", 1)]),
        ("Keq = 0,05", [lambda: setattr(rxn, "FixedEquilibriumConstant", S2_KEQ),
                        lambda: setattr(rxn, "Keq", S2_KEQ)]),
    ):
        try:
            attempt(desc, *funcs)
        except StepFailed:
            log(f"    → задайте вручную: Properties → Reactions → Rxn-1: {desc}")
    attempt("Rxn-1 → Set-1", lambda: rset.Reactions.Add(rxn), lambda: rset.AddReaction(rxn))
    attempt("Set-1 → пакет (Add to FP)", lambda: fp.ReactionSets.Add(rset),
            lambda: rset.AttachToFluidPackage(fp))
    return rset


def s2_build(app):
    log("\n=== S2: синтез аммиака (задача 4.1) ===")
    case = new_case(app)
    hold(case)
    fp = setup_basis(case, S2_COMPONENTS, "PR-S2")
    rset = None
    try:
        rset = s2_reaction(case, fp)
    except StepFailed:
        log("  ! Реакцию через COM создать не удалось. Создайте её вручную по таблице 5 "
            "инструкции (раздел 3.1), сохраните файл и запустите: python hysys_auto.py read s2 файл.hsc")
    fs = case.Flowsheet
    feed, vap, liq = add_stream(fs, "Feed1"), add_stream(fs, "Vap1"), add_stream(fs, "Liq1")
    q = add_energy(fs, "Q-100")
    names = fp_component_names(fp, S2_COMPONENTS)
    flows = [S2_FEED_FLOWS[S2_COMPONENTS.index(n)] if n in S2_COMPONENTS else 0.0 for n in names]
    tot = sum(flows)
    set_var(feed, "Temperature", S2_FEED["T_C"], "C")
    set_var(feed, "Pressure", S2_FEED["P_kPa"], "kPa")
    set_var(feed, "MolarFlow", tot, "kgmole/h")
    set_composition(feed, [f / tot for f in flows])
    rx = add_op(fs, "ERV-100", "equilibrium")
    connect_separator(rx, feed, vap, liq)
    attempt("ERV-100: энергия", lambda: setattr(rx, "EnergyStream", q))
    if rset is not None:
        try:
            attempt("ERV-100: Reaction Set", lambda: setattr(rx, "ReactionSet", rset),
                    lambda: setattr(rx, "ReactionSetName", "Set-1"))
        except StepFailed:
            log("    → вручную: ERV-100 → Reactions → Details → Reaction Set = Set-1")
    set_var(vap, "Temperature", S2_FEED["T_C"], "C")  # изотермический режим
    solve(case)
    save(case, "S2_ammonia.hsc")
    return case


def s2_read(case):
    fs = case.Flowsheet
    vap, liq = get_stream(fs, "Vap1"), get_stream(fs, "Liq1")
    try:
        names = fp_component_names(case.Flowsheet.FluidPackage, S2_COMPONENTS)
    except Exception:  # noqa: BLE001
        names = S2_COMPONENTS
    F = flow(vap)
    y = read_fractions(vap, names)
    out = {"Ftot": F, "Liq1": flow(liq)}
    if F is not None and y:
        nN2, nH2, nNH3 = (F * y.get(c, 0.0) for c in S2_COMPONENTS)
        out.update(nN2=nN2, nH2=nH2, nNH3=nNH3, yNH3=y.get("Ammonia"),
                   XN2=100 * (1.1 - nN2) / 1.1, extent=nNH3 - 0.2)
    return out


# ----------------------------------------------------------------- сценарий 3
def s3_build(app):
    log("\n=== S3: реактор Гиббса, этан (задача 4.2а) ===")
    case = new_case(app)
    hold(case)
    fp = setup_basis(case, S3_COMPONENTS, "PR-S3")
    fs = case.Flowsheet
    feed, vap, liq = add_stream(fs, "Feed"), add_stream(fs, "Vap1"), add_stream(fs, "Liq1")
    q = add_energy(fs, "Q-100")
    names = fp_component_names(fp, S3_COMPONENTS)
    set_var(feed, "Temperature", S3_FEED["T_C"], "C")
    set_var(feed, "Pressure", S3_FEED["P_kPa"], "kPa")
    set_var(feed, "MolarFlow", S3_FEED["F_kmolh"], "kgmole/h")
    set_composition(feed, [1.0 if n == "Ethane" else 0.0 for n in names])
    rx = add_op(fs, "GBR-100", "gibbs")
    connect_separator(rx, feed, vap, liq)
    attempt("GBR-100: энергия", lambda: setattr(rx, "EnergyStream", q))
    set_var(vap, "Temperature", S3_FEED["T_C"], "C")
    solve(case)
    save(case, "S3_gibbs.hsc")
    return case


def s3_read(case):
    fs = case.Flowsheet
    vap = get_stream(fs, "Vap1")
    try:
        names = fp_component_names(case.Flowsheet.FluidPackage, S3_COMPONENTS)
    except Exception:  # noqa: BLE001
        names = S3_COMPONENTS
    F = flow(vap)
    y = read_fractions(vap, names)
    out = {"Ftot": F}
    if y:
        e, en, h = y.get("Ethane", 0), y.get("Ethylene", 0), y.get("Hydrogen", 0)
        out.update(yC2H6=e, yC2H4=en, yH2=h, K=(en * h / e) if e > 0 else None)
        if F is not None:
            out["x"] = 1 - F * e / S3_FEED["F_kmolh"]
    return out


# ----------------------------------------------------------------- вывод
def fmt(v, nd=4):
    if v is None:
        return "—"
    return f"{v:.{nd}f}".replace(".", ",")


def report_table(title, rows, res, with_pr=False):
    log(f"\n{title}")
    hdr = f"{'Показатель':36s} {'Отчёт':>10s}" + (f" {'PR':>10s}" if with_pr else "") + f" {'HYSYS':>10s}"
    log(hdr)
    out = []
    for row in rows:
        if with_pr:
            name, key, rep, pr = row
        else:
            name, key, rep = row
            pr = None
        h = res.get(key)
        line = f"{name:36s} {fmt(rep):>10s}" + (f" {fmt(pr):>10s}" if with_pr else "") + f" {fmt(h):>10s}"
        log(line)
        out.append((name, rep, pr, h))
    return out


def write_excel(tables, study):
    try:
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        log("! openpyxl не установлен (pip install openpyxl) – Excel не записан")
        return
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Описание"
    ws0.append(["Перекрёстная проверка МХТП (вариант 8) в Aspen HYSYS", ""])
    ws0.append(["Дата", _dt.datetime.now().strftime("%d.%m.%Y %H:%M")])
    ws0.append(["Источник «Отчёт»", "МХТП_Отчет_7сем_Суханов_вар8.docx (корреляция McWilliams / Kp = 0,05 / ΔG° IFP)"])
    ws0.append(["Источник «PR»", "07_МХТП/src/pr_flash.py: Пенг–Робинсон, kij = 0 (ориентир для HYSYS)"])
    ws0.append(["Отклонение", "(HYSYS − Отчёт)/Отчёт, %; для S1 расхождение с отчётом ожидаемо (см. раздел 2.4 инструкции)"])
    ws0["A1"].font = Font(bold=True, size=13)
    ws0.column_dimensions["A"].width = 22
    ws0.column_dimensions["B"].width = 100
    fill = PatternFill("solid", fgColor="E2EFDA")
    for title, rows, with_pr in tables:
        ws = wb.create_sheet(title[:31])
        head = ["Показатель", "Отчёт"] + (["Ориентир PR"] if with_pr else []) + ["HYSYS", "Отклонение от отчёта, %"]
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        for name, rep, pr, h in rows:
            dev = None
            if h is not None and rep not in (None, 0):
                dev = round(100 * (h - rep) / rep, 3)
            row = [name, rep] + ([pr] if with_pr else []) + [h, dev]
            ws.append(row)
            ws.cell(ws.max_row, len(head) - 1).fill = fill
        ws.column_dimensions["A"].width = 38
        for col in "BCDE":
            ws.column_dimensions[col].width = 16
    if study:
        ws = wb.create_sheet("S1 исследование P2")
        ws.append(["P2 = P3, МПа", "e2 (Feed2)", "e3 (Feed3)", "Gas3, kgmole/h", "y CH4 (Gas3)", "y C3H8 (Gas3)"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for P, r in study:
            ws.append([P / 1000, r["e2"], r["e3"], r["Gas3"], r["yCH4_Gas3"], r["yC3H8_Gas3"]])
        ch = LineChart()
        ch.title = "HYSYS (PR): выход газа С3 от P2 = P3"
        ch.x_axis.title = "P2 = P3, МПа"
        ch.y_axis.title = "Gas3, kgmole/h (на 100 kgmole/h сырья)"
        ch.add_data(Reference(ws, min_col=4, min_row=1, max_row=ws.max_row), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        ws.add_chart(ch, "H2")
        for col in "ABCDEF":
            ws.column_dimensions[col].width = 16
    wb.save(XLSX)
    log(f"\n✓ Результаты: {XLSX}")


# ----------------------------------------------------------------- диагностика
def diag(app, progid):
    out = HERE / "hysys_diag.txt"
    lines = [f"ProgID: {progid}", f"Python: {sys.version}"]
    for attr in ("Version", "Name", "FullName"):
        try:
            lines.append(f"app.{attr} = {getattr(app, attr)}")
        except Exception as e:  # noqa: BLE001
            lines.append(f"app.{attr}: {e!r}")
    try:
        lines.append(f"Открытых файлов: {app.SimulationCases.Count}")
        if app.SimulationCases.Count:
            case = app.ActiveDocument if hasattr(app, "ActiveDocument") else app.SimulationCases.Item(0)
            fs = case.Flowsheet
            lines.append("Операции активного файла (имя → TypeName):")
            for i in range(fs.Operations.Count):
                op = fs.Operations.Item(i)
                lines.append(f"   {op.Name} → {op.TypeName}")
            lines.append("Потоки: " + ", ".join(str(fs.MaterialStreams.Item(i).Name)
                                                for i in range(fs.MaterialStreams.Count)))
            try:
                fp = case.Flowsheet.FluidPackage
                lines.append(f"Пакет: {fp.PropertyPackageName}; компоненты: {fp_component_names(fp, [])}")
            except Exception as e:  # noqa: BLE001
                lines.append(f"Пакет: {e!r}")
    except Exception as e:  # noqa: BLE001
        lines.append(f"Диагностика файла: {e!r}")
    out.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    print(f"\n✓ Записано: {out}")


# ----------------------------------------------------------------- main
def main(argv):
    if not argv or argv[0] in ("-h", "--help", "help"):
        print(__doc__)
        return 0
    cmd = argv[0].lower()
    dry = cmd == "selftest"
    if dry:
        sys.path.insert(0, str(HERE))
        argv = ["all"]
        cmd = "all"
    LOG.write_text(f"Журнал запуска {_dt.datetime.now():%d.%m.%Y %H:%M}: {' '.join(argv)}\n", encoding="utf-8")
    app, progid = connect(dry)
    log(f"Связь с HYSYS: {progid}")
    if cmd == "diag":
        diag(app, progid)
        return 0
    tables, study = [], []
    try:
        if cmd == "read":
            if len(argv) < 3:
                print("Использование: python hysys_auto.py read s1|s2|s3 файл.hsc")
                return 2
            which, path = argv[1].lower(), Path(argv[2]).resolve()
            case = open_case(app, path)
            if which == "s1":
                res, _ = s1_run(app, case, built=False)
                mode = list(res)[0]
                tables.append(("S1 (файл)", report_table("S1 – файл как есть (сверять с таблицей 3 или 4)",
                                                        S1_EXPECT["base"], res[mode], True), True))
            elif which == "s2":
                solve(case)
                tables.append(("S2 аммиак", report_table("S2 – синтез аммиака", S2_EXPECT, s2_read(case)), False))
            elif which == "s3":
                solve(case)
                tables.append(("S3 Гиббс", report_table("S3 – реактор Гиббса", S3_EXPECT, s3_read(case)), False))
        else:
            todo = ["s1", "s2", "s3"] if cmd == "all" else [cmd]
            for which in todo:
                try:
                    if which == "s1":
                        res, study = s1_run(app)
                        for mode, title in (("base", "S1 база 9,8-8,7 МПа"), ("6MPa", "S1 6,0 МПа")):
                            tables.append((title, report_table(f"S1 – {title}", S1_EXPECT[mode], res[mode], True), True))
                    elif which == "s2":
                        case = s2_build(app)
                        tables.append(("S2 аммиак", report_table("S2 – синтез аммиака", S2_EXPECT, s2_read(case)), False))
                    elif which == "s3":
                        case = s3_build(app)
                        tables.append(("S3 Гиббс", report_table("S3 – реактор Гиббса", S3_EXPECT, s3_read(case)), False))
                    else:
                        print(f"Неизвестная команда: {which}")
                        return 2
                except StepFailed as e:
                    log(f"\n✗ Сценарий {which.upper()} остановлен на шаге «{e}».")
                    log("  Что делать: соберите этот шаг вручную по инструкции, сохраните файл и "
                        f"запустите `python hysys_auto.py read {which} файл.hsc`; "
                        "пришлите hysys_log.txt и hysys_diag.txt для доработки скрипта.")
    except Exception:  # noqa: BLE001
        log("\n✗ Непредвиденная ошибка:\n" + traceback.format_exc())
    if tables:
        write_excel(tables, study)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
